#!/usr/bin/env python3
"""
build_leadership_feed.py — strip internal-only sheets before publishing.

WHY THIS EXISTS
The leadership dashboard is served from GitHub Pages, which is PUBLIC and
search-indexable (CLAUDE.md Hard Rule #7 and the owner's 2026-06-05 exception,
which warns about exactly this). The workbook carries the operational record, so
anything left in the file that the Pages build copies is published with it.

The Opportunity_Detail sheet added 2026-08-28 holds BD posture, win themes,
teaming status, compliance gaps, and staffing gaps for every pursuit. That is
internal planning material for the operations and delivery teams — it is not
leadership-summary content and must not be search-indexable. This script removes
it from the copy that gets published. The committed workbook is never modified.

The dashboard degrades cleanly without the sheet: its parser returns [] for a
missing sheet and every detail panel reads "Not yet recorded", so the published
leadership view keeps working exactly as it did before the sheet existed.

WHY THIS EDITS THE ZIP INSTEAD OF USING OPENPYXL
An .xlsx is a zip of XML parts. Removing a sheet with openpyxl means load + save,
and openpyxl does not preserve CACHED FORMULA RESULTS — it writes the formula
string but drops the last-computed value. The dashboard reads only cached values
(SheetJS `cell.v`), and nine of the eleven cells it pulls with getCell() are
formulas: Targets!C11-C14 (confirmed revenue, run-rate P/L, break-even gap, ARR
gap) and Targets!D34-D38 (the five BD monthly targets). An openpyxl round-trip on
the publish path would blank all nine and silently drop those KPIs to the
hardcoded defaults in the HTML.

So this operates on the package directly: it removes the sheet's parts and the
references to them, and copies every other part through byte-for-byte. Cached
values, styles, and everything else survive exactly as the source had them.

Safe to run repeatedly; it always writes a fresh copy from the source workbook.

    python scripts/build_leadership_feed.py [in.xlsx] [out.xlsx]

Defaults: VBX_Command_Center_v1.1.xlsx -> data/leadership_feed.xlsx

NOTE ON SCOPE: this strips the internal detail sheet only. The rest of the
workbook — partner, NDA, and revenue detail — is published as before under the
owner's standing exception. Widening the redaction is an owner decision, not a
default; add sheets to STRIP_SHEETS (or column-level redaction) when authorized.
"""
from __future__ import annotations

import posixpath
import re
import shutil
import sys
import zipfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pipeline_feed as pf  # noqa: E402

# Sheets removed from the published copy. Nothing in the workbook references
# these, so removing them cannot orphan a formula.
STRIP_SHEETS = (pf.DETAIL_SHEET,)

WORKBOOK_PART = "xl/workbook.xml"
WORKBOOK_RELS = "xl/_rels/workbook.xml.rels"
CONTENT_TYPES = "[Content_Types].xml"
# calcChain caches Excel's formula evaluation ORDER by sheet index. Removing a
# sheet invalidates those indices, so the part is dropped; Excel rebuilds it on
# next open and SheetJS never reads it. Cached VALUES live in the sheet parts and
# are untouched by this.
CALC_CHAIN = "xl/calcChain.xml"


def _resolve(target: str, base_part: str) -> str:
    """Resolve a relationship Target to a package part name."""
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _rels_part_for(part: str) -> str:
    d, name = posixpath.split(part)
    return posixpath.join(d, "_rels", name + ".rels")


def _referenced_parts(zf: zipfile.ZipFile, rels_part: str) -> set[str]:
    """Parts targeted by a .rels file (external targets ignored)."""
    if rels_part not in zf.namelist():
        return set()
    xml = zf.read(rels_part).decode("utf-8")
    out = set()
    for m in re.finditer(r"<Relationship\b[^>]*>", xml):
        tag = m.group(0)
        if 'TargetMode="External"' in tag:
            continue
        t = re.search(r'Target="([^"]+)"', tag)
        if t:
            out.add(_resolve(t.group(1), rels_part.replace("/_rels/", "/").replace(".rels", "")))
    return out


def strip_sheets(src: Path, dst: Path, sheet_names) -> list[str]:
    """Remove sheets from the package, copying all other parts verbatim.

    Returns the list of sheet names actually removed.
    """
    removed: list[str] = []
    with zipfile.ZipFile(src) as zin:
        names = zin.namelist()
        wb_xml = zin.read(WORKBOOK_PART).decode("utf-8")
        rels_xml = zin.read(WORKBOOK_RELS).decode("utf-8")

        # rId -> part name, from the workbook's relationships
        rid_to_part = {}
        for m in re.finditer(r"<Relationship\b[^>]*>", rels_xml):
            tag = m.group(0)
            rid = re.search(r'Id="([^"]+)"', tag)
            tgt = re.search(r'Target="([^"]+)"', tag)
            if rid and tgt:
                rid_to_part[rid.group(1)] = _resolve(tgt.group(1), WORKBOOK_PART)

        drop_parts: set[str] = set()
        drop_rids: set[str] = set()
        for m in re.finditer(r"<sheet\b[^>]*/>", wb_xml):
            tag = m.group(0)
            name = re.search(r'name="([^"]+)"', tag)
            rid = re.search(r'r:id="([^"]+)"', tag)
            if not name or name.group(1) not in sheet_names:
                continue
            removed.append(name.group(1))
            wb_xml = wb_xml.replace(tag, "")
            if rid:
                drop_rids.add(rid.group(1))
                part = rid_to_part.get(rid.group(1))
                if part:
                    drop_parts.add(part)
                    sheet_rels = _rels_part_for(part)
                    # Drop the sheet's own rels, plus any part ONLY it referenced
                    # (drawings, comments, printerSettings) so nothing leaks.
                    own = _referenced_parts(zin, sheet_rels)
                    others: set[str] = set()
                    for rp in names:
                        if rp.endswith(".rels") and rp != sheet_rels:
                            others |= _referenced_parts(zin, rp)
                    drop_parts |= {p for p in own if p not in others}
                    if sheet_rels in names:
                        drop_parts.add(sheet_rels)

        if not removed:
            shutil.copyfile(src, dst)
            return removed

        # definedNames pointing at a stripped sheet would dangle.
        for m in re.finditer(r"<definedName\b[^>]*>.*?</definedName>", wb_xml, re.S):
            body = m.group(0)
            if any(re.search(rf"(^|[^A-Za-z0-9_]){re.escape(s)}!", body) for s in removed):
                wb_xml = wb_xml.replace(body, "")

        for rid in drop_rids:
            rels_xml = re.sub(rf'<Relationship\b[^>]*Id="{re.escape(rid)}"[^>]*/>', "", rels_xml)
            rels_xml = re.sub(rf'<Relationship\b[^>]*Id="{re.escape(rid)}"[^>]*>.*?</Relationship>',
                              "", rels_xml, flags=re.S)

        if CALC_CHAIN in names:
            drop_parts.add(CALC_CHAIN)
            calc_rid = next((r for r, p in rid_to_part.items() if p == CALC_CHAIN), None)
            if calc_rid:
                rels_xml = re.sub(rf'<Relationship\b[^>]*Id="{re.escape(calc_rid)}"[^>]*/>', "", rels_xml)

        ct_xml = zin.read(CONTENT_TYPES).decode("utf-8")
        for part in drop_parts:
            ct_xml = re.sub(rf'<Override\b[^>]*PartName="/{re.escape(part)}"[^>]*/>', "", ct_xml)

        dst.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                if item.filename in drop_parts:
                    continue
                if item.filename == WORKBOOK_PART:
                    data = wb_xml.encode("utf-8")
                elif item.filename == WORKBOOK_RELS:
                    data = rels_xml.encode("utf-8")
                elif item.filename == CONTENT_TYPES:
                    data = ct_xml.encode("utf-8")
                else:
                    data = zin.read(item.filename)  # everything else, verbatim
                zout.writestr(item, data)
    return removed


def main() -> int:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src = Path(args[0] if len(args) > 0 else "VBX_Command_Center_v1.1.xlsx")
    dst = Path(args[1] if len(args) > 1 else "data/leadership_feed.xlsx")

    if not src.exists():
        print(f"ERROR: source workbook not found: {src}")
        return 1

    # Read-only inspection: a formula pointing at a stripped sheet would become
    # #REF! once published. Refuse rather than publish a broken workbook.
    wb = openpyxl.load_workbook(src, data_only=False)
    before = list(wb.sheetnames)
    dangling = sorted(
        cell
        for cell in pf.formula_set(wb)
        for sheet in STRIP_SHEETS
        if sheet in str(wb[cell.split("!")[0]][cell.split("!")[1]].value)
    )
    if dangling:
        print("ABORT: formulas reference a sheet slated for stripping:", dangling[:10])
        return 2

    removed = strip_sheets(src, dst, set(STRIP_SHEETS))

    # Verify the result opens and lost exactly what we intended.
    out = openpyxl.load_workbook(dst, data_only=False)
    expected = [s for s in before if s not in removed]
    if out.sheetnames != expected:
        print(f"ABORT: sheet list changed unexpectedly.\n  got:      {out.sheetnames}\n  expected: {expected}")
        return 2

    print(f"Leadership feed -> {dst}")
    print(f"  stripped: {', '.join(removed) if removed else 'nothing (sheet not present)'}")
    print(f"  kept {len(out.sheetnames)} sheet(s): {', '.join(out.sheetnames)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
