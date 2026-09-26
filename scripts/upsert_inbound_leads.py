#!/usr/bin/env python3
"""
upsert_inbound_leads.py — merge a partner-tracker inbound handoff into the workbook.

WHY THIS EXISTS
Inbound partner replies arrive as a tracker export (one record per firm). They
are UNSCREENED, so they must not land on Screenings, NDAs or Meetings: the BD
KPIs count every dated row on those sheets (CLAUDE.md §3). This script puts them
on the internal-only Inbound_Leads sheet instead, and applies the handoff's
partner-ecosystem NDA rows and open items in the same pass.

THIS FILE HOLDS NO DATA (Hard Rule #1)
Every firm, contact, date and action comes from the two input files below.
Keep them in data/_local_* (gitignored) or outside the repo; never commit them.

  firms.json   — the tracker export, verbatim:
      {"firms": [{"ID": "IN-XX-01", "Firm (legal name)": "...", "UEI": "...", ...}]}
      Keys are the tracker's column names (COLUMNS below). A key absent from a
      record leaves that cell blank — nothing is derived or backfilled.

  extras.json  — the parts of the handoff that are not per-firm tracker columns:
      {
        "schedule":       {"IN-XX-01": ["YYYY-MM-DD", "1:00 PM"]},
        "ecosystem_ndas": [{"id": "IN-XX-02", "doc": "<document id>",
                            "status": "<NDAs!F text>"}],
        "actions":        [{"action": "...", "owner": "...", "deadline": "YYYY-MM-DD" | null,
                            "success": "...", "status": "...", "source": "..."}]
      }
      Actions is PUBLISHED IN FULL — cite firms there by ID, never by name, or
      the leadership-feed leak scan fails the deploy.

RULES THE SCRIPT ENFORCES
  * Upsert: a firm already in the partner ecosystem (Screenings / NDAs /
    Agreements / Meetings) matches on UEI first, then legal name, and is kept
    OFF Inbound_Leads. Otherwise it matches an existing Inbound_Leads row on UEI,
    then legal name, and is updated in place. Re-runs never duplicate.
  * Classification, Score and SAM/SBA Verified are written EMPTY on insert and
    never touched on update: they fill by hand once a firm is actually screened.
    Call Date / Time update only for firms this handoff's schedule names.
  * An ecosystem NDA row is inserted once, with NO Date. NDAs!D is the Effective
    Date, which does not exist until the last signature; a Date here would count
    it as executed. A row already present is never rewritten, so the Date,
    status and classification entered by hand when it executes survive re-runs.
  * The 254-formula set must be unchanged, or nothing is saved.

    python scripts/upsert_inbound_leads.py FIRMS.json EXTRAS.json [--check] [--workbook PATH]

  --check   print what would change and write nothing.
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pipeline_feed as pf  # noqa: E402

# === CONSTANTS === #
DEFAULT_WORKBOOK = Path("VBX_Command_Center_v1.1.xlsx")
SHEET = "Inbound_Leads"
STYLE_SOURCE = "Actions"          # title / banner / header styling is copied from here
FIRST_DATA_ROW = pf.HEADER_ROW + 1

INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")               # blue = input
OVERDUE_FONT = Font(name="Arial", size=10, bold=True, color="FFDC2626")  # brand RED
OVERDUE_STATUS = "OVERDUE"
DATE_FORMAT = "mm/dd/yyyy"

TITLE = "  INBOUND LEADS  —  pre-Wave 1 partner replies (UNSCREENED)"
BANNER = ("  INTERNAL ONLY — stripped from the published copy. CLAIMED = firm's own "
          "unverified assertion. Classification / Score / Verified stay EMPTY until "
          "screened. Log a call in Meetings only after it happens.")

# (header, source key, width). Key None = schema-only column, always written EMPTY.
# "_call_date" / "_call_time" come from extras["schedule"], not the tracker export.
COLUMNS = [
    ("ID", "ID", 11),
    ("Firm (legal name)", "Firm (legal name)", 34),
    ("DBA", "DBA", 26),
    ("Lane", "Lane", 34),
    ("City", "City", 18),
    ("Contact", "Contact", 28),
    ("Email", "Email", 32),
    ("Phone", "Phone", 16),
    ("Website", "Website", 26),
    ("UEI", "UEI", 16),
    ("CAGE", "CAGE", 9),
    ("Primary NAICS", "Primary NAICS", 22),
    ("Why this firm", "Why this firm", 44),
    ("Status", "Status", 15),
    ("Call Date", "_call_date", 12),
    ("Call Time (CT)", "_call_time", 13),
    ("Response notes", "Response notes", 80),
    ("NDA status", "NDA status", 11),
    ("Next step", "Next step", 60),
    ("Owner", "Owner", 10),
    ("Source chat", "Source chat", 16),
    ("Reserve Pool / Wave 1 match", "Reserve Pool / Wave 1 match", 44),
    ("Classification", None, 15),
    ("Score", None, 8),
    ("SAM/SBA Verified", None, 16),
]
WRAP_COLUMNS = {"Lane", "Why this firm", "Response notes", "Next step",
                "Reserve Pool / Wave 1 match"}
CALL_KEYS = {"_call_date", "_call_time"}
COL_UEI = 1 + [h for h, _, _ in COLUMNS].index("UEI")
COL_NAME = 1 + [h for h, _, _ in COLUMNS].index("Firm (legal name)")

# Partner-ecosystem sheets: (sheet, name column, columns whose text may embed a UEI)
ECOSYSTEM = (("Screenings", 2, (4,)), ("NDAs", 2, ()), ("Agreements", 2, ()),
             ("Meetings", 2, ()))

NDA_NOTES_COL = 8                 # NDAs!H — withheld by the publish allowlist
NDA_DOC_COL = 5                   # NDAs!E
NDA_LAST_STYLED_COL = 7


# === HELPERS === #
def norm(name) -> str:
    """Case/punctuation-insensitive legal-name key."""
    return " ".join(str(name or "").lower().replace(",", " ").replace(".", " ").split())


def iso_date(value):
    return dt.date.fromisoformat(value) if value else None


def copy_style(dst, src) -> None:
    dst.font = copy.copy(src.font)
    dst.fill = copy.copy(src.fill)
    dst.alignment = copy.copy(src.alignment)
    dst.border = copy.copy(src.border)


def last_data_row(ws) -> int:
    rows = [r for r in range(FIRST_DATA_ROW, ws.max_row + 1) if ws.cell(r, 1).value]
    return rows[-1] if rows else pf.HEADER_ROW


# === INBOUND_LEADS SHEET === #
def ensure_sheet(wb):
    """Return Inbound_Leads, creating it (title, banner, row-4 headers) if absent."""
    if SHEET in wb.sheetnames:
        return wb[SHEET]
    src = wb[STYLE_SOURCE]
    ws = wb.create_sheet(SHEET)
    ws["A1"], ws["A2"] = TITLE, BANNER
    for col in range(1, len(COLUMNS) + 1):
        copy_style(ws.cell(1, col), src["A1" if col == 1 else "B1"])
        copy_style(ws.cell(2, col), src["A2" if col == 1 else "B2"])
    ws.row_dimensions[1].height = src.row_dimensions[1].height
    for col, (header, _, width) in enumerate(COLUMNS, start=1):
        copy_style(ws.cell(pf.HEADER_ROW, col, header), src["B4"])
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "C5"
    return ws


def ecosystem_match(wb, firm) -> str | None:
    """Where the firm is already tracked as a partner: UEI first, then legal name."""
    uei, name = firm.get("UEI"), norm(firm["Firm (legal name)"])
    for sheet, name_col, uei_cols in ECOSYSTEM:
        ws = wb[sheet]
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            if uei and any(uei in str(ws.cell(r, c).value or "") for c in uei_cols):
                return f"{sheet}!row {r} (UEI)"
            if name and norm(ws.cell(r, name_col).value) == name:
                return f"{sheet}!row {r} (legal name)"
    return None


def upsert_firms(wb, ws, firms, schedule) -> dict:
    by_uei, by_name = {}, {}
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if ws.cell(r, 1).value:
            if ws.cell(r, COL_UEI).value:
                by_uei[ws.cell(r, COL_UEI).value] = r
            by_name[norm(ws.cell(r, COL_NAME).value)] = r
    next_row = last_data_row(ws) + 1

    log = {"inserted": [], "updated": [], "ecosystem": {}}
    for firm in firms:
        where = ecosystem_match(wb, firm)
        if where:
            log["ecosystem"][firm["ID"]] = where
            continue
        row = (by_uei.get(firm["UEI"]) if firm.get("UEI") else None) \
            or by_name.get(norm(firm["Firm (legal name)"]))
        if row:
            log["updated"].append(firm["ID"])
        else:
            row, next_row = next_row, next_row + 1
            log["inserted"].append(firm["ID"])
        is_new = firm["ID"] in log["inserted"]
        record = dict(firm)
        scheduled = firm["ID"] in schedule
        if scheduled:
            call_date, call_time = schedule[firm["ID"]]
            record["_call_date"], record["_call_time"] = iso_date(call_date), call_time
        for col, (header, key, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row, col)
            # On an update, never touch what the export does not own: the
            # screening columns (filled by hand once screened) and a call slot
            # this handoff's schedule does not mention.
            if not is_new and (key is None or (key in CALL_KEYS and not scheduled)):
                continue
            cell.value = record.get(key) if key else None   # absent stays absent
            cell.font = copy.copy(INPUT_FONT)
            cell.alignment = Alignment(wrap_text=header in WRAP_COLUMNS, vertical="top")
            if header == "Call Date":
                cell.number_format = DATE_FORMAT
    return log


# === PARTNER-ECOSYSTEM NDA ROWS === #
def upsert_ecosystem_ndas(wb, firms_by_id, entries, ecosystem) -> list[int]:
    ws = wb["NDAs"]
    if ws.cell(pf.HEADER_ROW, NDA_NOTES_COL).value is None:
        copy_style(ws.cell(pf.HEADER_ROW, NDA_NOTES_COL, "Notes"),
                   ws.cell(pf.HEADER_ROW, NDA_LAST_STYLED_COL))
        ws.column_dimensions[get_column_letter(NDA_NOTES_COL)].width = 80
    rows = []
    for entry in entries:
        if entry["id"] not in ecosystem:
            raise SystemExit(f"ABORT: {entry['id']} has an ecosystem NDA but is not "
                             "tracked in the partner ecosystem — refusing to split it.")
        firm = firms_by_id[entry["id"]]
        last = last_data_row(ws)
        existing = next((r for r in range(FIRST_DATA_ROW, last + 1)
                         if ws.cell(r, NDA_DOC_COL).value == entry["doc"]), None)
        if existing:
            # Insert-once. After this, NDAs is the record: the Effective Date,
            # "Executed" and the classification are entered by hand when the
            # NDA executes, and a re-run must never revert them.
            rows.append(existing)
            continue
        row = last + 1
        # D (Date) and G (Classification) start empty: not executed until the
        # last signature, not classified until screened.
        values = [ws.cell(last, 1).value + 1, firm["Firm (legal name)"], firm.get("Contact"),
                  None, entry["doc"], entry["status"], None, firm.get("Response notes")]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col)
            cell.value = value
            copy_style(cell, ws.cell(last, min(col, NDA_LAST_STYLED_COL)))
            if col == NDA_NOTES_COL:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        rows.append(row)
    return rows


# === ACTIONS === #
def append_actions(wb, actions) -> list[int]:
    ws = wb["Actions"]
    have = {ws.cell(r, 2).value for r in range(FIRST_DATA_ROW, ws.max_row + 1)}
    row, added = last_data_row(ws), []
    for a in actions:
        if a["action"] in have:
            continue
        row += 1
        values = [ws.cell(row - 1, 1).value + 1, a["action"], a["owner"],
                  iso_date(a.get("deadline")), a["success"], a["status"], a["source"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            copy_style(cell, ws.cell(FIRST_DATA_ROW, col))
            cell.number_format = ws.cell(FIRST_DATA_ROW, col).number_format
            if col == 6 and a["status"] == OVERDUE_STATUS:
                cell.font = copy.copy(OVERDUE_FONT)
        added.append(row)
    return added


# === MAIN === #
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    ap.add_argument("firms")
    ap.add_argument("extras")
    ap.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    ap.add_argument("--check", action="store_true", help="dry run; write nothing")
    args = ap.parse_args()

    firms = json.loads(Path(args.firms).read_text())["firms"]
    extras = json.loads(Path(args.extras).read_text())
    ids = [f["ID"] for f in firms]
    if len(ids) != len(set(ids)):
        raise SystemExit("ABORT: duplicate IDs in the firms file.")

    wb = openpyxl.load_workbook(args.workbook)
    before = pf.formula_set(wb)
    ws = ensure_sheet(wb)
    log = upsert_firms(wb, ws, firms, extras.get("schedule", {}))
    nda_rows = upsert_ecosystem_ndas(wb, dict(zip(ids, firms)),
                                     extras.get("ecosystem_ndas", []), log["ecosystem"])
    action_rows = append_actions(wb, extras.get("actions", []))

    if pf.formula_set(wb) != before:
        raise SystemExit("ABORT: formula set changed — nothing saved.")

    print(f"Inbound_Leads: {len(log['inserted'])} inserted, {len(log['updated'])} updated")
    for fid, where in log["ecosystem"].items():
        print(f"  {fid} already in the partner ecosystem ({where}) — kept off Inbound_Leads")
    print(f"NDAs ecosystem rows (inserted or already present, untouched): {nda_rows or 'none'}"
          f"  |  Actions rows added: {action_rows or 'none'}")
    print(f"Formulas: {len(before)} (unchanged)")
    if args.check:
        print("--check: nothing written.")
        return 0
    wb.save(args.workbook)
    print(f"Saved {args.workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
