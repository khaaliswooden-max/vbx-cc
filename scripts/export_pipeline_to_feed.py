#!/usr/bin/env python3
"""
export_pipeline_to_feed.py — one-way export: workbook -> feed.

Primarily the one-time MIGRATION that seeds data/pipeline_feed.jsonl from the
current committed workbook. Also usable as a rescue tool to rebuild the feed
from the workbook if they ever drift. It never touches the workbook.

Reads BOTH the Pipeline sheet and the Opportunity_Detail sheet (BD posture +
technical scope), so a rescue rebuild does not silently drop the detail fields.
Author-chosen `key` values in the existing feed are preserved too — the feed is
matched on OPP-ID, and only rows with no existing line fall back to key == opp_id.

Usage:
    python scripts/export_pipeline_to_feed.py \
        [VBX_Command_Center_v1.1.xlsx] [data/pipeline_feed.jsonl]
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pipeline_feed as pf  # noqa: E402

WORKBOOK = sys.argv[1] if len(sys.argv) > 1 else "VBX_Command_Center_v1.1.xlsx"
FEED = sys.argv[2] if len(sys.argv) > 2 else "data/pipeline_feed.jsonl"


def read_detail(wb) -> dict[str, dict]:
    """Opportunity_Detail rows keyed by Opp ID ({} when the sheet is absent)."""
    if pf.DETAIL_SHEET not in wb.sheetnames:
        return {}
    ws = wb[pf.DETAIL_SHEET]
    out: dict[str, dict] = {}
    for r in range(pf.FIRST_DATA_ROW, ws.max_row + 1):
        vals = {name: ws.cell(r, c).value for c, name in enumerate(pf.DETAIL_COLUMNS, 1)}
        opp_id = vals.get("opp_id")
        if not opp_id:
            continue
        rec = {}
        for field in pf.DETAIL_FIELDS:
            v = vals.get(field)
            if v in (None, ""):
                continue
            if field in pf.LIST_FIELDS:
                rec[field] = pf.as_list(v)
            elif field in pf.DETAIL_DATE_FIELDS:
                rec[field] = pf.iso_date(v)
            else:
                rec[field] = v
        out[opp_id] = rec
    return out


def main() -> int:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    ws = wb[pf.SHEET]
    detail = read_detail(wb)
    # Keep author-chosen keys from the feed we are about to overwrite.
    existing_keys = {r["opp_id"]: r["key"] for r in pf.load_feed(FEED)
                     if r.get("opp_id") and r.get("key")}

    last = pf.HEADER_ROW
    for r in range(pf.FIRST_DATA_ROW, ws.max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, pf.LAST_COL + 1)):
            last = r

    records: list[dict] = []
    for r in range(pf.FIRST_DATA_ROW, last + 1):
        vals = {name: ws.cell(r, c).value for c, name in enumerate(pf.COLUMNS, 1)}
        opp_id = vals["opp_id"]
        rec = {
            # migrated rows: key == frozen OPP-ID, unless the feed already
            # carries an author-chosen key for this opportunity.
            "key": existing_keys.get(opp_id, opp_id),
            "opp_id": opp_id,
            "name": vals["name"],
            "client": vals["client"],
            "naics": vals["naics"],
            "stage": vals["stage"],
            "value": vals["value"] if isinstance(vals["value"], int) else None,
            "date_identified": pf.iso_date(vals["date_identified"]),
            "stage_date": pf.iso_date(vals["stage_date"]),
            "owner": vals["owner"],
            "notes": vals["notes"],
        }
        rec.update(detail.get(opp_id, {}))
        records.append(rec)

    problems = pf.validate(records)
    if problems:
        print("FEED VALIDATION FAILED:")
        for p in problems:
            print("  -", p)
        return 1

    Path(FEED).parent.mkdir(parents=True, exist_ok=True)
    pf.dump_feed(FEED, records)
    detailed = sum(1 for r in records if pf.has_detail(r))
    print(f"Exported {len(records)} Pipeline rows -> {FEED} "
          f"({detailed} with BD/technical detail from {pf.DETAIL_SHEET})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
